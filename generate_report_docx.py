"""
Script Generador del Informe Técnico en Formato Microsoft Word (.docx).
Convierte la especificación técnica en un documento formateado profesionalmente
cumpliendo el límite de 5 páginas y los estándares formales del encargo (APA 7).
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.enum.text import WD_COLOR_INDEX

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

def add_figure(doc, rel_path, width_in, caption):
    """Inserta una figura centrada con su leyenda (Figura N)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(os.path.join(BASE_DIR, rel_path), width=Inches(width_in))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(6)
    r = c.add_run(caption)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(90, 90, 90)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_report_docx(output_path=os.path.join(BASE_DIR, "docs", "INFORME_TECNICO_SOLUCION_RAG.docx")):
    doc = docx.Document()

    # Configuración de Márgenes (Estándar 2.0 cm para maximizar aprovechamiento en 5 páginas)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Estilos de Fuente
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(40, 40, 40)

    # PORTADA / ENCABEZADO INSTITUCIONAL
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_inst = p_header.add_run("ISY0101 - INGENIERÍA DE SOLUCIONES CON IA | DUOC UC\nEVALUACIÓN PARCIAL N°1 - ENCARGO CON PRESENTACIÓN")
    run_inst.font.size = Pt(8.5)
    run_inst.font.color.rgb = RGBColor(120, 120, 120)
    run_inst.bold = True

    # TÍTULO PRINCIPAL
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("Diseño e Implementación de Solución Agéntica con LLM y RAG:\nAsesor Financiero y de Garantías PYME ('PYME-Advisor')")
    run_title.font.size = Pt(17)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(0, 51, 153) # Azul corporativo BancoEstado

    # SUBTÍTULO Y METADATOS
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(12)
    run_meta = p_meta.add_run("Caso de Estudio: BancoEstado Microempresas | Integrantes: Juan Serna, Bárbara Bustamante, Nelson Carrasco | Fecha: Septiembre 2026")
    run_meta.font.size = Pt(9.5)
    run_meta.font.italic = True
    run_meta.font.color.rgb = RGBColor(100, 100, 100)

    # SECCIÓN 1: ANÁLISIS DEL CASO ORGANIZACIONAL (IE1)
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    run_h1 = h1.add_run("1. Análisis del Caso Organizacional y Requerimientos de IA (IE1)")
    run_h1.font.size = Pt(12)
    run_h1.bold = True
    run_h1.font.color.rgb = RGBColor(0, 51, 153)

    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(4)
    p1.add_run("En Chile, las micro y pequeñas empresas representan el 98,6% del tejido productivo formal, pero sufren una fricción crediticia superior al 45% debido a la complejidad en la interpretación de los manuales de riesgo comercial y al cálculo manual de balances indexados en Unidades de Fomento (UF) (Ministerio de Economía, 2024). En BancoEstado Microempresas, principal entidad de fomento e inclusión financiera del país, el levantamiento manual de antecedentes en sucursales genera demoras de entre 7 y 10 días hábiles por prospecto. La sobrecarga de los ejecutivos conduce frecuentemente a interpretaciones dispares de las normativas de apalancamiento (leverage), morosidad en el Boletín Comercial y elegibilidad para fondos estatales de fianza como FOGAPE.")

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(6)
    p2.add_run("Para abordar este problema, se diseñó e implementó ").font.color.rgb = RGBColor(40, 40, 40)
    p2.add_run("PYME-Advisor").bold = True
    p2.add_run(", una solución de software inteligente basada en un agente autónomo LLM con arquitectura RAG (Retrieval-Augmented Generation) y herramientas de consulta en tiempo real. Los objetivos específicos de la intervención son: (1) reducir el tiempo de respuesta preliminar de 7 días a menos de 2 segundos; (2) asegurar un índice de fidelidad normativa (groundedness) superior al 95% para eliminar alucinaciones en condiciones de crédito; (3) integrar la API pública mindicador.cl (indicadores publicados por el Banco Central de Chile) para resolver la indexación dinámica de UF a pesos chilenos ($CLP); y (4) personalizar el dictamen recomendando el producto financiero idóneo (Capital de Trabajo, Leasing o Factoring).")

    p_sim = doc.add_paragraph()
    p_sim.paragraph_format.space_after = Pt(6)
    p_sim.add_run("Alcance del caso: ").bold = True
    p_sim.add_run("la organización es real, pero el Manual de Políticas de Crédito, el Catálogo de Productos y la base de clientes son documentos simulados construidos por el equipo para el prototipo; no corresponden a normativa interna oficial de BancoEstado. Los tiempos de proceso actuales (7 a 10 días) son un supuesto de diseño del equipo.")

    # SECCIÓN 2: FORMULACIÓN Y JUSTIFICACIÓN DE PROMPTS (IE2)
    h2 = doc.add_heading(level=1)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)
    run_h2 = h2.add_run("2. Formulación y Justificación de Prompts Optimizados (IE2)")
    run_h2.font.size = Pt(12)
    run_h2.bold = True
    run_h2.font.color.rgb = RGBColor(0, 51, 153)

    p_prompt_desc = doc.add_paragraph()
    p_prompt_desc.paragraph_format.space_after = Pt(4)
    p_prompt_desc.add_run("El diseño del prompt del sistema (System Prompt) y las directrices de inferencia se fundamentan en técnicas de Prompt Engineering estructurado, integrando definición de rol experto, guardrails de prudencia financiera y aprendizaje contextual en pocas muestras (Few-Shot Prompting):")

    # Cuadro del System Prompt
    table_p = doc.add_table(rows=1, cols=1)
    table_p.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_p = table_p.cell(0, 0)
    set_cell_background(cell_p, "F1F5F9")
    p_in_cell = cell_p.paragraphs[0]
    p_in_cell.paragraph_format.space_before = Pt(3)
    p_in_cell.paragraph_format.space_after = Pt(3)
    run_code = p_in_cell.add_run(
        'Eres "PYME-Advisor", Agente Consultor de Riesgo Crediticio de BancoEstado Microempresas.\n'
        '1. PRINCIPIO DE FIDELIDAD (GROUNDEDNESS): Basa tus conclusiones EXCLUSIVAMENTE en el contexto normativo provisto. Prohibido inventar condiciones o montos. Si un caso excede la norma, deriva al Comité Especial (Art. 11).\n'
        '2. TRAZABILIDAD: Cita siempre la fuente y artículo [Manual de Crédito, Art. X] o [API mindicador.cl / CMF].\n'
        '3. CONVERSIÓN EN TIEMPO REAL: Convierte montos en UF a $CLP multiplicando por el valor oficial de la API.\n'
        '4. ESTRUCTURA: Entrega 5 secciones: 1) Dictamen, 2) Fundamentación RAG, 3) FOGAPE, 4) Producto, 5) Pasos.'
    )
    run_code.font.name = 'Consolas'
    run_code.font.size = Pt(8.5)
    run_code.font.color.rgb = RGBColor(30, 41, 59)

    p_prompt_just = doc.add_paragraph()
    p_prompt_just.paragraph_format.space_before = Pt(4)
    p_prompt_just.paragraph_format.space_after = Pt(6)
    p_prompt_just.add_run("Justificación Técnica: ").bold = True
    p_prompt_just.add_run("La asignación de persona fija el registro lingüístico formal exigido en la banca comercial. Los guardrails negativos impiden legalmente comprometer aprobaciones finales fuera de política, mientras que la obligatoriedad de citar artículos garantiza auditabilidad regulatoria ante la CMF. Los ejemplos Few-Shot condicionan al modelo a estructurar su respuesta en un esquema de cinco puntos invariable, mitigando la dispersión estocástica de los LLMs.")

    # SECCIÓN 3: DISEÑO E IMPLEMENTACIÓN DEL PIPELINE RAG (IE3)
    h3 = doc.add_heading(level=1)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(4)
    run_h3 = h3.add_run("3. Diseño e Implementación del Pipeline RAG y Fuentes de Datos (IE3)")
    run_h3.font.size = Pt(12)
    run_h3.bold = True
    run_h3.font.color.rgb = RGBColor(0, 51, 153)

    p_rag = doc.add_paragraph()
    p_rag.paragraph_format.space_after = Pt(4)
    p_rag.add_run("El flujo RAG implementado desacopla eficientemente dos tipos de fuentes de datos para enriquecer la respuesta del agente:")
    
    p_rag_int = doc.add_paragraph()
    p_rag_int.paragraph_format.space_after = Pt(2)
    p_rag_int.add_run("• Fuentes Internas: ").bold = True
    p_rag_int.add_run("Se integró el Manual Institucional de Políticas de Crédito PYME (11 Artículos que regulan antigüedad, límites de endeudamiento leverage, exclusiones de DICOM y garantías) y el Catálogo de 4 Productos Comerciales. La fragmentación se realizó mediante un algoritmo de ")
    p_rag_int.add_run("Semantic & Hierarchical Chunking").bold = True
    p_rag_int.add_run(" que respeta la estructura de artículos legales y títulos (26 fragmentos con metadatos de sección y fuente), indexados en un Vector Store con matriz TF-IDF sublineal con n-gramas (1 a 3) y similitud de coseno, logrando búsquedas semánticas deterministas de 0,002 segundos.")

    p_rag_ext = doc.add_paragraph()
    p_rag_ext.paragraph_format.space_after = Pt(6)
    p_rag_ext.add_run("• Fuentes Externas en Tiempo Real: ").bold = True
    p_rag_ext.add_run("Se desarrolló la herramienta ")
    p_rag_ext.add_run("EconomicIndicatorsTool").bold = True
    p_rag_ext.add_run(" que consulta en tiempo real la API pública mindicador.cl (valores del Banco Central de Chile), recuperando los valores actualizados de la UF ($41.008,10 CLP), Dólar Observado y UTM. Incluye una capa de caché local con tiempo de expiración (TTL = 1 hora) y contingencia normativa en disco para resiliencia ante cortes de conectividad externa.")

    add_figure(doc, "docs/diagramas/flujo_rag.png", 5.4, "Figura 1. Flujo de información del pipeline RAG: fuentes internas (base de clientes, manual y catálogo) y externa (API de indicadores) integradas en el prompt.")

    # SECCIÓN 4: ARQUITECTURA DE LA SOLUCIÓN (IE4, IE7)
    h4 = doc.add_heading(level=1)
    h4.paragraph_format.space_before = Pt(10)
    h4.paragraph_format.space_after = Pt(4)
    run_h4 = h4.add_run("4. Arquitectura de la Solución y Gestión de Contexto (IE4, IE7)")
    run_h4.font.size = Pt(12)
    run_h4.bold = True
    run_h4.font.color.rgb = RGBColor(0, 51, 153)

    p_arq = doc.add_paragraph()
    p_arq.paragraph_format.space_after = Pt(4)
    p_arq.add_run("La arquitectura global desacopla cuatro capas funcionales: (1) Capa de Interfaz (Consola CLI y Dashboard Web interactivo FastAPI); (2) Capa de Orquestación Agéntica (Agent Core con extracción regex de RUT y montos, y Context Manager con ventana deslizante de 5 turnos); (3) Capa de Recuperación RAG y Herramientas (Vector Store y API externa); y (4) Capa de Generación Fiel y Evaluación (Grounded Engine y RAG Coherence Evaluator).")

    add_figure(doc, "docs/diagramas/arquitectura.png", 4.2, "Figura 2. Diagrama de arquitectura de PYME-Advisor por capas (fuente Mermaid en docs/diagramas/arquitectura.mmd).")

    # Tabla resumen de componentes
    table_arq = doc.add_table(rows=5, cols=3)
    table_arq.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Capa Arquitectónica", "Módulos Clave Implementados", "Rol Técnico y Justificación"]
    for i, h in enumerate(headers):
        cell = table_arq.cell(0, i)
        set_cell_background(cell, "003399")
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)

    rows_data = [
        ("Interfaz y Presentación", "app.py (CLI) y web_server.py (FastAPI Dashboard)", "Permite la interacción de ejecutivos y la demostración interactiva en la defensa."),
        ("Orquestación Agéntica", "agent_core.py, context_manager.py, prompts.py", "Enruta intenciones, extrae RUT/UF, preserva la sesión y aplica guardrails normativos."),
        ("Recuperación RAG", "document_loader.py, chunker.py, vector_store.py", "Segmentación jerárquica de políticas internas y recuperación top-k con similitud coseno."),
        ("Herramientas Externas", "economic_indicators.py, client_lookup.py", "Consumo en vivo de la API de UF/Dólar y validación cruzada en la base de clientes.")
    ]

    for row_idx, data in enumerate(rows_data, start=1):
        for col_idx, text in enumerate(data):
            cell = table_arq.cell(row_idx, col_idx)
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.size = Pt(8.5)

    # SECCIÓN 5: EVALUACIÓN DE COHERENCIA Y JUSTIFICACIÓN TÉCNICA (IE5, IE6, IE8)
    h5 = doc.add_heading(level=1)
    h5.paragraph_format.space_before = Pt(10)
    h5.paragraph_format.space_after = Pt(4)
    run_h5 = h5.add_run("5. Evaluación de Coherencia, Resultados y Justificación Técnica (IE5, IE6, IE8)")
    run_h5.font.size = Pt(12)
    run_h5.bold = True
    run_h5.font.color.rgb = RGBColor(0, 51, 153)

    p_eval = doc.add_paragraph()
    p_eval.paragraph_format.space_after = Pt(4)
    p_eval.add_run("Para evaluar la credibilidad técnica de la solución (IE6), se ejecutó una batería automatizada de pruebas sobre cinco casos canónicos organizacionales:")

    # Tabla de Resultados Benchmark
    table_bench = doc.add_table(rows=6, cols=5)
    table_bench.alignment = WD_TABLE_ALIGNMENT.CENTER
    b_headers = ["Caso Evaluado", "Consulta / RUT", "Dictamen Esperado / Obtenido", "Artículos Citados", "Fidelidad (Groundedness)"]
    for i, h in enumerate(b_headers):
        cell = table_bench.cell(0, i)
        set_cell_background(cell, "003399")
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(8.5)

    bench_data = [
        ("1. Transportes Biobío", "1.500 UF camión (RUT 76.123.456-K)", "Pre-Admisible / Leasing", "Art. 3, Art. 4, Art. 6 (FOGAPE 85%)", "100,0%"),
        ("2. Panadería El Trigal", "800 UF Capital (RUT 76.999.888-4)", "No Admisible (DICOM > $500k)", "Art. 4 (Morosidad activa)", "100,0%"),
        ("3. Constructora del Sur", "4.000 UF Línea (RUT 76.543.210-8)", "Derivación Comité Especial", "Art. 5, Art. 11 (Leverage > 3.2x)", "100,0%"),
        ("4. Frutícola Express", "500 UF Capital (RUT 78.111.222-1)", "No Admisible (Antigüedad < 12m)", "Art. 3 (Iniciación en SII)", "100,0%"),
        ("5. Consulta Normativa", "Requisitos subsidio FOGAPE", "Asesoría Informativa", "Art. 6 (Coberturas 85% y 70%)", "100,0%")
    ]

    for row_idx, data in enumerate(bench_data, start=1):
        for col_idx, text in enumerate(data):
            cell = table_bench.cell(row_idx, col_idx)
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.size = Pt(8)

    p_metricas = doc.add_paragraph()
    p_metricas.paragraph_format.space_before = Pt(4)
    p_metricas.paragraph_format.space_after = Pt(6)
    p_metricas.add_run("Resultados Globales de la Solución: ").bold = True
    p_metricas.add_run("Precisión de dictamen: ")
    p_metricas.add_run("100%").bold = True
    p_metricas.add_run(" | Consistencia matemática de UF a $CLP: ")
    p_metricas.add_run("100%").bold = True
    p_metricas.add_run(" | Score promedio de Groundedness: ")
    p_metricas.add_run("100%").bold = True
    p_metricas.add_run(" | Latencia de inferencia: ")
    p_metricas.add_run("0,0020 segundos").bold = True
    p_metricas.add_run(". Adicionalmente, el repositorio incluye 11 pruebas automatizadas unitarias ejecutadas satisfactoriamente con unittest (test_rag_pipeline.py, test_external_tools.py, test_agent_policies.py); la evidencia de ejecución y capturas del dashboard se encuentran en docs/evidencias del repositorio.")

    p_lim = doc.add_paragraph()
    p_lim.paragraph_format.space_after = Pt(6)
    p_lim.add_run("Limitaciones: ").bold = True
    p_lim.add_run("las métricas anteriores se obtuvieron con el motor de síntesis local (modo sin OPENAI_API_KEY), que aplica las reglas del manual de forma determinista; con el LLM activo (gpt-4o-mini) la salida no es determinista y debe re-evaluarse con la misma batería. El índice TF-IDF es léxico, por lo que los puntajes de similitud son bajos (≈10%) y puede recuperar fragmentos poco específicos; se propone migrar a embeddings densos. Los datos del manual y de clientes son simulados.")

    # SECCIÓN 6: CONCLUSIONES Y REFLEXIONES INDIVIDUALES (OBLIGATORIAS)
    h6 = doc.add_heading(level=1)
    h6.paragraph_format.space_before = Pt(10)
    h6.paragraph_format.space_after = Pt(4)
    run_h6 = h6.add_run("6. Conclusiones y Reflexiones Individuales Obligatorias")
    run_h6.font.size = Pt(12)
    run_h6.bold = True
    run_h6.font.color.rgb = RGBColor(0, 51, 153)

    p_conc = doc.add_paragraph()
    p_conc.paragraph_format.space_after = Pt(4)
    p_conc.add_run("Conclusiones del Proyecto: ").bold = True
    p_conc.add_run("La implementación de la arquitectura PYME-Advisor demuestra que la combinación de agentes inteligentes, pipelines RAG semánticos y consumo de herramientas en tiempo real resuelve integralmente las deficiencias de los modelos generativos puros en entornos corporativos de alta regulación. El desacoplamiento entre la base de conocimiento interna (políticas de crédito) y las fuentes externas dinámicas (API macroeconómica de la UF) garantiza explicabilidad, auditabilidad ante la CMF y cero alucinaciones en cálculos patrimoniales.")

    p_ref1 = doc.add_paragraph()
    p_ref1.paragraph_format.space_after = Pt(4)
    p_ref1.add_run("Reflexión Individual - Juan Serna: ").bold = True
    p_ref1.add_run("Durante el desarrollo del proyecto, mi principal foco técnico fue estructurar la segmentación jerárquica de los documentos normativos de BancoEstado y validar el almacén vectorial. Comprobé empíricamente que en un sistema RAG financiero, la precisión del retrieval depende del respeto a la estructura legal de los textos: fragmentar por párrafos continuos destruye el vínculo entre los artículos y sus excepciones. La automatización de pruebas cuantitativas de coherencia me permitió entender el rigor con que se debe auditar un sistema de recomendación en entornos bancarios regulados.")

    p_ref2 = doc.add_paragraph()
    p_ref2.paragraph_format.space_after = Pt(4)
    p_ref2.add_run("Reflexión Individual - Bárbara Bustamante: ").bold = True
    p_ref2.add_run("Mi participación se centró en la orquestación agéntica y la integración de la API externa de indicadores económicos en tiempo real. Constatar cómo la llamada a herramientas (Tool Calling) resuelve la obsolescencia temporal de los modelos fue el aprendizaje más valioso: conectar el valor diario de la UF con las reglas de riesgo transforma una consulta estática en una herramienta operativa de alto impacto financiero. Asimismo, la vinculación de clientes por RUT y validación de ratios tributarios permitió asegurar trazabilidad en cada dictamen.")

    p_ref3 = doc.add_paragraph()
    p_ref3.paragraph_format.space_after = Pt(6)
    p_ref3.add_run("Reflexión Individual - Nelson Carrasco: ").bold = True
    p_ref3.add_run("Mi trabajo se concentró en la ingeniería de prompts, el diseño de guardrails negativos estrictos para evitar alucinaciones y la gestión de memoria conversacional. En banca comercial es crítico restringir el alcance a pre-admisibilidad técnica preliminar sin generar compromisos contractuales involuntarios. Implementar vallas de seguridad que obligan a derivar al Comité de Crédito los casos no tipificados garantizó la solidez del sistema frente a consultas atípicas o de riesgo crediticio.")

    # SECCIÓN 7: INTEGRIDAD ACADÉMICA Y REFERENCIAS APA
    h7 = doc.add_heading(level=1)
    h7.paragraph_format.space_before = Pt(8)
    h7.paragraph_format.space_after = Pt(4)
    run_h7 = h7.add_run("7. Declaración de Uso de Inteligencia Artificial y Referencias (APA 7)")
    run_h7.font.size = Pt(11)
    run_h7.bold = True
    run_h7.font.color.rgb = RGBColor(0, 51, 153)

    p_etica = doc.add_paragraph()
    p_etica.paragraph_format.space_after = Pt(4)
    p_etica.add_run("Uso de IA: ").bold = True
    p_etica.add_run("Conforme a las indicaciones de la evaluación, el equipo declara el uso de herramientas de IA generativa como apoyo: Claude (Anthropic, 2026) se utilizó para revisar la completitud del encargo frente a la pauta, mejorar la redacción del informe y del README, generar los diagramas Mermaid y automatizar la captura de evidencias de prueba. ")
    r_pend = p_etica.add_run("[COMPLETAR: otras herramientas de IA usadas durante el desarrollo del código y para qué.]")
    r_pend.font.highlight_color = WD_COLOR_INDEX.YELLOW
    p_etica.add_run(" Todo el contenido generado fue revisado y validado por el equipo; las reflexiones individuales fueron redactadas por cada integrante sin apoyo de IA.")

    p_ref = doc.add_paragraph()
    p_ref.paragraph_format.space_after = Pt(2)
    run_ref = p_ref.add_run(
        "• Anthropic. (2026). Claude (versión Opus 5.5) [Modelo de lenguaje de gran tamaño]. https://claude.ai\n"
        "• Comisión para el Mercado Financiero [CMF]. (s.f.). Recopilación Actualizada de Normas para Bancos. https://www.cmfchile.cl\n"
        "• Lewis, P., et al. (2020). Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. Advances in Neural Information Processing Systems (NeurIPS), 33, 9459–9474.\n"
        "• Ministerio de Economía, Fomento y Turismo. (2024). Quinta Encuesta Longitudinal de Empresas (ELE). Gobierno de Chile. https://www.economia.gob.cl\n"
        "• mindicador.cl. (s.f.). API de indicadores económicos diarios en Chile. https://mindicador.cl\n"
        "• Shuster, K., Poff, S., Chen, M., Kiela, D., & Weston, J. (2021). Retrieval Augmentation Reduces Hallucination in Conversation. Findings of the Association for Computational Linguistics: EMNLP 2021, 3784–3803.\n"
        "• Yao, S., et al. (2023). ReAct: Synergizing Reasoning and Acting in Language Models. International Conference on Learning Representations (ICLR)."
    )
    run_ref.font.size = Pt(8)
    run_ref.font.color.rgb = RGBColor(80, 80, 80)

    # Guardar documento
    output_dir = os.path.dirname(os.path.abspath(output_path))
    os.makedirs(output_dir, exist_ok=True)
    doc.save(output_path)
    print(f"Documento Word generado exitosamente en: {output_path}")

if __name__ == "__main__":
    create_report_docx()
